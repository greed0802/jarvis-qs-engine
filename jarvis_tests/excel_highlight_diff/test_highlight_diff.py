from openpyxl import load_workbook
from engine.highlight_engine import compare_maps
from output.highlight_writer import highlight_excel

def build_map(ws):
    data = {}
    for row in range(2, ws.max_row + 1):
        item = ws[f"A{row}"].value
        data[item] = {
            "description": ws[f"B{row}"].value,
            "qty": ws[f"C{row}"].value
        }
    return data


def test_highlight_diff():
    old_wb = load_workbook("old.xlsx")
    new_wb = load_workbook("new.xlsx")

    old_ws = old_wb.active
    new_ws = new_wb.active

    old_map = build_map(old_ws)
    new_map = build_map(new_ws)

    diff = compare_maps(old_map, new_map)

    highlight_excel("new.xlsx", diff, "jarvis-tests/excel_highlight_diff/reports/output_highlight.xlsx")

    assert diff is not None