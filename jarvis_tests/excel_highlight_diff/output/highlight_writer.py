from openpyxl import load_workbook
from jarvis_tests.excel_highlight_diff.styles.styles import RED, GREEN, YELLOW, BLUE

def highlight_excel(file_path, diff, output_path):
    wb = load_workbook(file_path)

    ws = wb.active

    # assume headers: item = col A, desc = B, qty = C

    for row in range(2, ws.max_row + 1):
        item = ws[f"A{row}"].value

        if item in diff["missing"]:
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = RED

        if item in diff["added"]:
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = GREEN

        if item in diff["qty_changes"]:
            ws[f"C{row}"].fill = YELLOW

        if item in diff["desc_changes"]:
            ws[f"B{row}"].fill = BLUE

    wb.save(output_path)