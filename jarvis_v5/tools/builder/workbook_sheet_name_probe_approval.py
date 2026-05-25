from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from jarvis_v5.config import ATTACHMENTS_DIR

NO_ENGINE_SAFETY = {
    "workbook_read": False,
    "engine_called": False,
    "excel_created": False,
    "contract_only": True,
    "legacy_builder_called": False,
}

ALLOWED_PROBE_FIELDS = ["sheet_names", "sheet_count"]

BLOCKED_OPERATIONS = [
    "read_cells",
    "read_formulas",
    "read_workbook_dimensions",
    "parse_workbook",
    "call_builder_engine",
    "call_legacy_builder",
    "create_excel",
    "upgrade_preview_readiness",
    "upgrade_export_readiness",
]

ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}
FORBIDDEN_EXACT_RAW_PATH_KEYS = {
    "saved_path",
    "absolute_path",
    "file_path",
    "local_path",
    "stored_path",
    "resolved_path",
    "raw_path",
    "attachment_path",
    "directory",
    "parent",
    "full_path",
    "path",
}


def _contract_payload(contract: Any | None) -> dict[str, Any] | None:
    if contract is None:
        return None
    if hasattr(contract, "model_dump"):
        payload = contract.model_dump(mode="json")
        return payload if isinstance(payload, dict) else None
    return contract if isinstance(contract, dict) else None


def _extension_from_filename(filename: Any | None) -> str | None:
    if not isinstance(filename, str) or not filename.strip() or "." not in filename:
        return None
    return "." + filename.strip().lower().rsplit(".", 1)[-1]


def _safe_attachment_file(workbook_ref: Any | None) -> tuple[Path | None, str | None]:
    if not isinstance(workbook_ref, dict) or not workbook_ref:
        return None, "workbook_ref_not_found"
    filename = workbook_ref.get("filename")
    extension = _extension_from_filename(filename)
    if extension not in ALLOWED_EXTENSIONS:
        return None, "workbook_extension_not_allowed"
    saved_path = workbook_ref.get("saved_path")
    if not isinstance(saved_path, str) or not saved_path.strip():
        return None, "workbook_saved_path_missing"
    try:
        candidate = Path(saved_path).expanduser().resolve()
        attachment_root = ATTACHMENTS_DIR.resolve()
        candidate.relative_to(attachment_root)
    except Exception:
        return None, "unsafe_workbook_path"
    if not candidate.exists() or not candidate.is_file():
        return None, "workbook_file_not_found"
    return candidate, None


def _contains_forbidden_exact_key(payload: Any) -> bool:
    if isinstance(payload, dict):
        return any(str(k) in FORBIDDEN_EXACT_RAW_PATH_KEYS or _contains_forbidden_exact_key(v) for k, v in payload.items())
    if isinstance(payload, list):
        return any(_contains_forbidden_exact_key(v) for v in payload)
    return False


def _base_safety(**overrides: Any) -> dict[str, Any]:
    safety = {
        **dict(NO_ENGINE_SAFETY),
        "workbook_opened": False,
        "workbook_closed": False,
        "limited_metadata_read": False,
        "sheet_names_read": False,
        "workbook_content_read": False,
        "workbook_parsed": False,
        "cells_read": False,
        "formulas_read": False,
        "raw_path_value_returned": False,
        "forbidden_path_key_returned": False,
        "preview_readiness_upgraded": False,
        "export_readiness_upgraded": False,
    }
    safety.update(overrides)
    return safety


def evaluate_workbook_sheet_name_probe_approval(
    *,
    conversation_id: str,
    client_event_id: str | None = None,
    contract: Any | None = None,
    approval_granted: bool = False,
) -> dict[str, Any]:
    """Approval-gated workbook sheet-name probe.

    This is the only owner of the alpha.34 sheet-name probe boundary. It opens
    an approved uploaded workbook only to read workbook-level sheet names, then
    closes the workbook immediately. It never reads workbook content, formulas,
    cells, dimensions, or output data and never mutates Builder state.
    """
    contract_payload = _contract_payload(contract)
    contract_found = contract_payload is not None
    workbook_ref = contract_payload.get("workbook_ref") if contract_payload else None
    workbook_ref_found = isinstance(workbook_ref, dict) and bool(workbook_ref)
    approval_required = not bool(approval_granted)
    blocked_by_policy = [
        "sheet_name_probe_contract",
        "cells_read_disabled",
        "formulas_read_disabled",
        "workbook_parse_disabled",
        "execution_kill_switch",
    ]
    sheet_names: list[str] = []
    workbook_opened = False
    workbook_closed = False
    limited_metadata_read = False
    sheet_names_read = False

    if not contract_found:
        blocked = True
        reason = "engine_contract_not_found"
        blocked_by_policy.append("engine_contract_not_found")
    elif not workbook_ref_found:
        blocked = True
        reason = "workbook_ref_not_found"
        blocked_by_policy.append("workbook_ref_not_found")
    elif approval_required:
        blocked = True
        reason = "sheet_name_probe_approval_required"
        blocked_by_policy.append("sheet_name_probe_approval_required")
    else:
        candidate, path_reason = _safe_attachment_file(workbook_ref)
        if candidate is None:
            blocked = True
            reason = path_reason or "workbook_not_safe_for_sheet_name_probe"
            blocked_by_policy.append(reason)
        else:
            workbook = None
            try:
                workbook = load_workbook(filename=candidate, read_only=True, data_only=False, keep_links=False)
                workbook_opened = True
                sheet_names = list(workbook.sheetnames)
                limited_metadata_read = True
                sheet_names_read = True
                blocked = False
                reason = "sheet_name_probe_approved_sheet_names_only"
            except Exception:
                blocked = True
                reason = "workbook_sheet_name_probe_failed"
                blocked_by_policy.append("workbook_sheet_name_probe_failed")
                sheet_names = []
                limited_metadata_read = False
                sheet_names_read = False
            finally:
                if workbook is not None:
                    try:
                        workbook.close()
                    finally:
                        workbook_closed = True

    safety = _base_safety(
        workbook_opened=workbook_opened,
        workbook_closed=workbook_closed,
        limited_metadata_read=limited_metadata_read,
        sheet_names_read=sheet_names_read,
    )
    response = {
        "conversation_id": conversation_id,
        "client_event_id": client_event_id,
        "contract_id": contract_payload.get("contract_id") if contract_payload else None,
        "route": "builder_workbook_sheet_name_probe_approval",
        "message": "Workbook sheet-name probe approval evaluated. No cells, formulas, dimensions, Builder engine, or Excel output were used.",
        "approval_required": approval_required,
        "approval_granted": bool(approval_granted),
        "blocked": blocked,
        "reason": reason,
        "contract_found": contract_found,
        "workbook_ref_found": workbook_ref_found,
        "workbook_opened": workbook_opened,
        "workbook_closed": workbook_closed,
        "limited_metadata_read": limited_metadata_read,
        "sheet_names_read": sheet_names_read,
        "sheet_names": sheet_names,
        "sheet_count": len(sheet_names),
        "workbook_read": False,
        "workbook_content_read": False,
        "workbook_parsed": False,
        "cells_read": False,
        "formulas_read": False,
        "engine_called": False,
        "excel_created": False,
        "contract_only": True,
        "legacy_builder_called": False,
        "raw_path_value_returned": False,
        "forbidden_path_key_returned": False,
        "preview_readiness_upgraded": False,
        "export_readiness_upgraded": False,
        "blocked_by_policy": blocked_by_policy,
        "allowed_probe_fields": list(ALLOWED_PROBE_FIELDS),
        "blocked_operations": list(BLOCKED_OPERATIONS),
        "safety": safety,
        "readiness": {
            "status": "sheet_name_probe_blocked" if blocked else "sheet_name_probe_complete",
            "preview_readiness_upgraded": False,
            "export_readiness_upgraded": False,
            "next_gate": "fix_contract_or_grant_sheet_name_probe_approval" if blocked else "future_workbook_read_policy_review",
            "blocked_by_policy": blocked_by_policy,
        },
        "next_safe_action": "Review sheet names only. Workbook cells, formulas, dimensions, parsing, Builder engine, and Excel output remain blocked.",
    }
    response["forbidden_path_key_returned"] = _contains_forbidden_exact_key(response)
    response["safety"]["forbidden_path_key_returned"] = response["forbidden_path_key_returned"]
    return response
